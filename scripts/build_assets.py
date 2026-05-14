#!/usr/bin/env python3
"""
Processa fotos brutas dos mentores e gera profile.webp + work.webp
em /public/mentors/{slug}/.

Para mentores que ainda NÃO têm foto, gera um placeholder visual com
nome + categoria + marca "FOTO PENDENTE" bem visível.

Uso típico:
    python build_assets.py

Sem argumentos, ele assume:
    - planilha:  ./Face-Talk-Mentors.xlsx  (ou ./data/mentors.xlsx)
    - origem:    ./photos/                  (com subpastas profile/ e work/)
    - destino:   ./public/mentors/

Pode sobrescrever:
    python build_assets.py --photos ~/Desktop/face-talk-photos \\
                           --xlsx data/mentors.xlsx \\
                           --out public/mentors

Modo "mesma imagem pra profile e work" (útil pra mockup):
    python build_assets.py --mirror

Para NÃO gerar placeholders (deixa pasta vazia):
    python build_assets.py --no-placeholders

Dependências:
    pip install pillow openpyxl
    pip install pillow-heif        # opcional, só se você tiver fotos .heic do iPhone
"""
from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from pathlib import Path

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:
    sys.exit("Faltou Pillow. Instala com:  pip install pillow")

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Faltou openpyxl. Instala com:  pip install openpyxl")

# Suporte opcional a HEIC (fotos do iPhone)
try:
    from pillow_heif import register_heif_opener
    register_heif_opener()
except ImportError:
    pass

# ---------- specs ----------
PROFILE_SIZE = (1024, 1024)
WORK_SIZE    = (1600, 900)
WEBP_QUALITY = 80
SUPPORTED_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".avif", ".heic", ".heif", ".bmp", ".tiff"}

# Cor de fundo do placeholder por categoria (gradient claro → escuro)
CATEGORY_COLORS = {
    "Carreira e Negócios":  ((30, 64, 175),   (15, 23, 92)),    # azul
    "Tecnologia":           ((88, 28, 135),   (49, 14, 76)),    # roxo
    "Saúde e Bem Estar":    ((6, 95, 70),     (3, 51, 38)),     # verde
    "Criatividade":         ((190, 24, 93),   (113, 14, 55)),   # rosa
    "Estilo de Vida":       ((194, 65, 12),   (113, 38, 7)),    # laranja
    "Gastronomia":          ((153, 27, 27),   (89, 16, 16)),    # vermelho
}
DEFAULT_COLORS = ((55, 65, 81), (24, 28, 38))  # cinza


# ---------- helpers ----------

def normalize(text: str) -> str:
    text = unicodedata.normalize("NFKD", text).encode("ASCII", "ignore").decode("ASCII")
    return re.sub(r"[^a-z0-9]", "", text.lower())


def smart_crop(img: Image.Image, target_w: int, target_h: int) -> Image.Image:
    src_w, src_h = img.size
    target_ratio = target_w / target_h
    src_ratio = src_w / src_h
    if src_ratio > target_ratio:
        new_w = int(src_h * target_ratio)
        left = (src_w - new_w) // 2
        img = img.crop((left, 0, left + new_w, src_h))
    elif src_ratio < target_ratio:
        new_h = int(src_w / target_ratio)
        top = (src_h - new_h) // 2
        img = img.crop((0, top, src_w, top + new_h))
    return img.resize((target_w, target_h), Image.LANCZOS)


def initials_from(name: str) -> str:
    parts = [p for p in name.split() if p and p[0].isalpha()]
    if not parts:
        return "?"
    if len(parts) == 1:
        return parts[0][:2].upper()
    return (parts[0][0] + parts[-1][0]).upper()


def load_font(size: int) -> ImageFont.ImageFont:
    """Tenta achar uma fonte do sistema; cai pra default se não."""
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",  # macOS
        "/System/Library/Fonts/Helvetica.ttc",                # macOS
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",  # Linux
        "C:\\Windows\\Fonts\\arialbd.ttf",                    # Windows
    ]
    for p in candidates:
        try:
            return ImageFont.truetype(p, size)
        except (OSError, IOError):
            continue
    return ImageFont.load_default()


def make_gradient(size: tuple[int, int], top_color, bottom_color) -> Image.Image:
    """Gradient vertical simples."""
    w, h = size
    base = Image.new("RGB", (w, h), top_color)
    top = Image.new("RGB", (1, h))
    for y in range(h):
        t = y / max(1, h - 1)
        r = int(top_color[0] * (1 - t) + bottom_color[0] * t)
        g = int(top_color[1] * (1 - t) + bottom_color[1] * t)
        b = int(top_color[2] * (1 - t) + bottom_color[2] * t)
        top.putpixel((0, y), (r, g, b))
    return top.resize((w, h))


def text_size(draw: ImageDraw.ImageDraw, text: str, font) -> tuple[int, int]:
    bbox = draw.textbbox((0, 0), text, font=font)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def make_placeholder(name: str, category: str, kind: str) -> Image.Image:
    """Gera placeholder bonito mas inequivocamente 'PENDENTE'."""
    size = PROFILE_SIZE if kind == "profile" else WORK_SIZE
    w, h = size
    top, bottom = CATEGORY_COLORS.get(category, DEFAULT_COLORS)
    img = make_gradient(size, top, bottom)
    draw = ImageDraw.Draw(img)

    if kind == "profile":
        # Iniciais grandes
        initials = initials_from(name)
        font_big = load_font(int(h * 0.42))
        tw, th = text_size(draw, initials, font_big)
        draw.text(((w - tw) / 2, (h - th) / 2 - 30), initials, fill="white", font=font_big)

        # "FOTO PENDENTE" embaixo
        font_tag = load_font(int(h * 0.04))
        tag = "FOTO PENDENTE"
        tw, th = text_size(draw, tag, font_tag)
        draw.text(((w - tw) / 2, h - th - 50), tag, fill=(255, 255, 255, 200), font=font_tag)

    else:  # work
        # Nome do mentor centralizado, com categoria abaixo
        font_name = load_font(int(h * 0.10))
        font_meta = load_font(int(h * 0.045))
        font_tag = load_font(int(h * 0.04))

        # Wrap simples se nome for muito longo
        max_w = int(w * 0.8)
        words = name.split()
        lines, cur = [], ""
        for word in words:
            test = (cur + " " + word).strip()
            tw, _ = text_size(draw, test, font_name)
            if tw <= max_w:
                cur = test
            else:
                if cur:
                    lines.append(cur)
                cur = word
        if cur:
            lines.append(cur)

        line_h = text_size(draw, "Hg", font_name)[1] + 10
        block_h = line_h * len(lines)
        y = (h - block_h) / 2 - 30
        for line in lines:
            tw, _ = text_size(draw, line, font_name)
            draw.text(((w - tw) / 2, y), line, fill="white", font=font_name)
            y += line_h

        # Categoria
        cat_label = category.upper()
        tw, th = text_size(draw, cat_label, font_meta)
        draw.text(((w - tw) / 2, y + 20), cat_label, fill=(255, 255, 255, 180), font=font_meta)

        # "FOTO PENDENTE" no canto inferior direito
        tag = "FOTO PENDENTE"
        tw, th = text_size(draw, tag, font_tag)
        draw.text((w - tw - 40, h - th - 40), tag, fill=(255, 255, 255, 200), font=font_tag)

    return img


def save_webp(img: Image.Image, dst: Path) -> int:
    dst.parent.mkdir(parents=True, exist_ok=True)
    img.save(dst, format="WEBP", quality=WEBP_QUALITY, method=6)
    return dst.stat().st_size


# ---------- planilha ----------

def load_mentors_from_xlsx(xlsx_path: Path) -> dict[str, dict]:
    """{slug_normalizado: {slug, name, category}}"""
    wb = load_workbook(xlsx_path)
    sh = wb["Página1"] if "Página1" in wb.sheetnames else wb.active
    headers = {c.value: c.column for c in sh[1]}
    for required in ("Name", "Slug", "Category"):
        if required not in headers:
            sys.exit(f"Planilha sem coluna '{required}'. Headers: {list(headers)}")

    out = {}
    for r in range(2, sh.max_row + 1):
        slug = sh.cell(row=r, column=headers["Slug"]).value
        if not slug:
            continue
        out[normalize(slug)] = {
            "slug": slug,
            "name": sh.cell(row=r, column=headers["Name"]).value,
            "category": sh.cell(row=r, column=headers["Category"]).value,
        }
    return out


# ---------- processamento ----------

def find_photos(folder: Path) -> dict[str, Path]:
    if not folder.exists():
        return {}
    out = {}
    for p in folder.iterdir():
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTS:
            stem = normalize(p.stem)
            if stem in out:
                print(f"  ⚠️  Conflito: {p.name} colide com {out[stem].name}")
                continue
            out[stem] = p
    return out


def process_real_photos(kind: str, photos_dir: Path, mentors: dict[str, dict],
                        out_root: Path, target: tuple[int, int]) -> dict:
    photos = find_photos(photos_dir)
    print(f"\n=== {kind.upper()} (fotos reais) ===")
    print(f"Origem: {photos_dir}")
    print(f"Encontradas: {len(photos)} imagens")

    matched = []
    unmatched = []
    has_real = set()

    for stem, src_path in sorted(photos.items()):
        if stem in mentors:
            slug = mentors[stem]["slug"]
            dst = out_root / slug / f"{kind}.webp"
            try:
                with Image.open(src_path) as img:
                    img = img.convert("RGB")
                    out = smart_crop(img, *target)
                size = save_webp(out, dst)
                matched.append((src_path.name, slug, size))
                has_real.add(slug)
                print(f"  ✓ {src_path.name:35} → {slug}/{kind}.webp ({size//1024} KB)")
            except Exception as e:
                print(f"  ✗ {src_path.name}: erro — {e}")
                unmatched.append(src_path.name)
        else:
            unmatched.append(src_path.name)
            print(f"  ? {src_path.name} — nenhum slug bateu")

    return {"matched": matched, "unmatched": unmatched, "has_real": has_real}


def fill_placeholders(kind: str, mentors: dict[str, dict], out_root: Path,
                      has_real: set[str]) -> int:
    """Gera placeholder pra todo mentor que não tem .webp real."""
    target = PROFILE_SIZE if kind == "profile" else WORK_SIZE
    count = 0
    print(f"\n=== {kind.upper()} (placeholders) ===")
    for m in mentors.values():
        if m["slug"] in has_real:
            continue
        dst = out_root / m["slug"] / f"{kind}.webp"
        # Não sobrescreve se já existe (caso usuário tenha rodado antes)
        if dst.exists():
            continue
        img = make_placeholder(m["name"], m["category"], kind)
        size = save_webp(img, dst)
        count += 1
    print(f"Gerados: {count} placeholders")
    return count


# ---------- main ----------

def find_default_xlsx() -> Path | None:
    for p in [Path("Face-Talk-Mentors.xlsx"), Path("data/mentors.xlsx"),
              Path("data/Face-Talk-Mentors.xlsx")]:
        if p.exists():
            return p
    return None


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--photos", default="photos")
    ap.add_argument("--xlsx", default=None)
    ap.add_argument("--out", default="public/mentors")
    ap.add_argument("--mirror", action="store_true",
                    help="Usa fotos da pasta profile/ pra gerar profile E work")
    ap.add_argument("--no-placeholders", action="store_true",
                    help="Não gera placeholders pra mentores sem foto")
    args = ap.parse_args()

    xlsx_path = Path(args.xlsx).expanduser().resolve() if args.xlsx else find_default_xlsx()
    if not xlsx_path or not xlsx_path.exists():
        sys.exit("Planilha não encontrada. Use --xlsx caminho/para/mentors.xlsx")
    photos = Path(args.photos).expanduser().resolve()
    out = Path(args.out).expanduser().resolve()

    print(f"Planilha: {xlsx_path}")
    print(f"Fotos:    {photos}{' (não existe — só placeholders)' if not photos.exists() else ''}")
    print(f"Destino:  {out}")
    print(f"Placeholders: {'desligado' if args.no_placeholders else 'ligado'}")

    mentors = load_mentors_from_xlsx(xlsx_path)
    print(f"Mentores na planilha: {len(mentors)}")

    profile_dir = photos / "profile"
    work_dir = photos / ("profile" if args.mirror else "work")

    p_real = process_real_photos("profile", profile_dir, mentors, out, PROFILE_SIZE) \
             if photos.exists() else {"matched": [], "unmatched": [], "has_real": set()}
    w_real = process_real_photos("work", work_dir, mentors, out, WORK_SIZE) \
             if photos.exists() else {"matched": [], "unmatched": [], "has_real": set()}

    p_placeholders = w_placeholders = 0
    if not args.no_placeholders:
        p_placeholders = fill_placeholders("profile", mentors, out, p_real["has_real"])
        w_placeholders = fill_placeholders("work", mentors, out, w_real["has_real"])

    print("\n" + "=" * 50)
    print("RESUMO")
    print("=" * 50)
    print(f"Profile: {len(p_real['matched']):>3} fotos reais  +  "
          f"{p_placeholders:>3} placeholders  =  "
          f"{len(p_real['matched']) + p_placeholders} arquivos")
    print(f"Work:    {len(w_real['matched']):>3} fotos reais  +  "
          f"{w_placeholders:>3} placeholders  =  "
          f"{len(w_real['matched']) + w_placeholders} arquivos")
    if p_real["unmatched"] or w_real["unmatched"]:
        print(f"\nFotos sem slug correspondente (verifique o nome):")
        for f in p_real["unmatched"] + w_real["unmatched"]:
            print(f"  - {f}")


if __name__ == "__main__":
    main()
